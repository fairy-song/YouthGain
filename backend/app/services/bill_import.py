"""Bounded CSV/XLSX bill adapters and account-scoped, atomic import batches.

The original file is never persisted. A committed batch is one learning-store
document so a retry cannot expose half of an import on any storage backend.
"""
import csv
import hashlib
import io
import json
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import chain

from openpyxl import load_workbook

from . import learning_store
from .checkin_service import BEIJING

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1500
MAX_XLSX_UNCOMPRESSED = 12 * 1024 * 1024
MAX_BATCH_BYTES = 850 * 1024
MAX_AMOUNT = Decimal('10000000.00')
PARSER_VERSION = 4
STANDARD_COLUMNS = {'date', 'type', 'amount'}
WECHAT_COLUMNS = {'交易时间', '交易类型', '交易对方', '收/支', '当前状态', '交易单号'}
VALID_TYPES = {'expense', 'refund', 'income', 'transfer'}
RECEIPT_KINDS = {'red_packet', 'external_transfer', 'other_income', 'internal_transfer'}


def _clean(value):
    return str(value or '').strip().lstrip('\ufeff').lstrip('`').strip()


def _amount(value):
    try:
        amount = Decimal(_clean(value).replace('¥', '').replace('￥', '').replace(',', ''))
        if not amount.is_finite() or not 0 < amount <= MAX_AMOUNT or amount.as_tuple().exponent < -2:
            raise InvalidOperation
        return str(amount.quantize(Decimal('0.01')))
    except (InvalidOperation, ValueError):
        raise ValueError('金额必须是大于 0、最多两位小数的人民币金额')


def _date(value):
    try:
        parsed = date.fromisoformat(_clean(value)[:10])
        if parsed > datetime.now(BEIJING).date():
            raise ValueError
        return parsed.isoformat()
    except ValueError:
        raise ValueError('日期无效或晚于今天')


def _xlsx_lines(content):
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ValueError('XLSX 文件格式无效，请重新导出微信账单')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (len(members) > 100 or sum(item.file_size for item in members) > MAX_XLSX_UNCOMPRESSED
                    or any(item.flag_bits & 1 for item in members)):
                raise ValueError('XLSX 文件过大或已加密，无法安全预览')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows()
                beginning = []
                for line_number, cells in enumerate(rows, start=1):
                    values = [cell.value for cell in cells]
                    beginning.append((line_number, values,
                                      any(getattr(cell, 'data_type', '') == 'f' for cell in cells)))
                    if line_number >= 21:
                        break
                if not any(_header(values) for _, values, _ in beginning):
                    continue
                for line_number, values, formula in chain(beginning,
                        ((number, [cell.value for cell in cells],
                          any(getattr(cell, 'data_type', '') == 'f' for cell in cells))
                         for number, cells in enumerate(rows, start=len(beginning) + 1))):
                    if line_number > MAX_ROWS + 25:
                        raise ValueError(f'单次最多导入 {MAX_ROWS} 行，请按日期拆分账单')
                    yield line_number, values, formula
                return
        finally:
            workbook.close()
    except (zipfile.BadZipFile, OSError, KeyError, ValueError) as error:
        if isinstance(error, ValueError) and str(error).startswith(('单次最多', 'XLSX 文件过大')):
            raise
        raise ValueError('XLSX 文件无法读取，请重新导出或检查文件是否损坏') from error
    raise ValueError('未识别到账单表头；请确认工作表包含微信支付交易明细')


def _header(line):
    columns = [_clean(cell) for cell in line]
    current = set(columns)
    if STANDARD_COLUMNS <= current:
        return columns, 'standard'
    if WECHAT_COLUMNS <= current and ('金额(元)' in current or '金额' in current):
        return columns, 'wechat'
    return None


def _read(content, hint='auto', file_type='csv'):
    if not content or len(content) > MAX_BYTES:
        raise ValueError('请选择不超过 2 MB 的 CSV 或 XLSX 账单')
    if hint not in ('auto', 'standard', 'wechat'):
        raise ValueError('账单来源无效')
    if file_type == 'xlsx':
        lines = _xlsx_lines(content)
    elif file_type == 'csv':
        if content.startswith(b'PK'):
            raise ValueError('文件内容是 XLSX，请使用 .xlsx 文件名重新选择')
        decoded = None
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None or '\x00' in decoded:
            raise ValueError('文件编码无法识别')
        try:
            reader = csv.reader(io.StringIO(decoded, newline=''), strict=True)
            lines = ((number, row, False) for number, row in enumerate(reader, start=1))
        except csv.Error:
            raise ValueError('CSV 格式不完整，请检查引号和换行')
    else:
        raise ValueError('只支持 CSV 或 XLSX 文件')
    found = None
    records = []
    try:
        for line_number, line, has_formula in lines:
            if line_number > MAX_ROWS + 25:
                raise ValueError(f'单次最多导入 {MAX_ROWS} 行，请按日期拆分账单')
            if found is None:
                header = _header(line) if line_number <= 21 else None
                if header:
                    columns, source = header
                    found = True
                    if hint != 'auto' and hint != source:
                        raise ValueError('所选账单来源与文件表头不符')
                    if len(columns) != len(set(columns)):
                        raise ValueError('账单表头含重复列，请核对文件')
                continue
            if not any(_clean(cell) for cell in line):
                continue
            if source == 'wechat' and _clean(line[0]).startswith(('总交易单数', '收入笔数', '支出笔数')):
                break
            if has_formula:
                records.append({'row_id': line_number, 'decision': 'review', 'reason': '含公式，请核对原账单'})
            elif len(line) != len(columns):
                records.append({'row_id': line_number, 'decision': 'review', 'reason': '列数与表头不一致'})
            else:
                row = dict(zip(columns, map(_clean, line)))
                records.append(_standard(row, line_number) if source == 'standard' else _wechat(row, line_number))
            if len(records) > MAX_ROWS:
                raise ValueError(f'单次最多导入 {MAX_ROWS} 行，请按日期拆分账单')
    except csv.Error as error:
        raise ValueError('CSV 格式不完整，请检查引号和换行') from error
    if found is None:
        raise ValueError('未识别到账单表头；请使用青盈标准模板或受支持的微信支付账单')
    if not records:
        raise ValueError('账单没有交易行')
    return source, records


def _common(line_number, kind, amount, spent, category, merchant, note, source_id, related_id, status):
    record = {'row_id': line_number, 'transaction_type': kind, 'amount': None, 'date': None,
              'category': _clean(category)[:50] or '未分类', 'merchant': _clean(merchant)[:100],
              'note': _clean(note)[:200], 'source_transaction_id': _clean(source_id)[:100],
              'related_source_id': _clean(related_id)[:100], 'merchant_order_id': '',
              'receipt_kind': None, 'direction': '', 'decision': 'import', 'reason': ''}
    try:
        record['amount'] = _amount(amount)
        record['date'] = _date(spent)
        if kind not in VALID_TYPES:
            raise ValueError('交易类型无法识别')
        if status not in ('completed', 'paid', 'success', '成功', '支付成功', '交易成功', '已完成',
                          '已收钱', '已到账', '退款成功', '已存入零钱', '已转入零钱通', '对方已收钱',
                          '已转账', '提现已到账', '资金已到账', '对方已退还'):
            raise ValueError('交易尚未确认完成')
    except ValueError as error:
        record.update(decision='review', reason=str(error))
    for field, limit in ((category, 50), (merchant, 100), (note, 200), (source_id, 100), (related_id, 100)):
        if len(_clean(field)) > limit:
            record.update(decision='review', reason='文本字段过长，请核对原账单')
            break
    return record


def _standard(row, line_number):
    currency = _clean(row.get('currency') or 'CNY').upper()
    record = _common(line_number, _clean(row.get('type')).lower(), row.get('amount'),
                     row.get('date'), row.get('category'), row.get('merchant'), row.get('note'),
                     row.get('transaction_id'), row.get('related_transaction_id'),
                     _clean(row.get('status') or 'completed').lower())
    if currency not in ('CNY', 'RMB', '人民币'):
        record.update(decision='review', reason='仅支持人民币账单')
    if not record['source_transaction_id']:
        record.update(decision='review', reason='缺少交易号，无法跨文件去重')
    record['merchant_order_id'] = _clean(row.get('order_id'))[:100]
    record['direction'] = 'income' if record['transaction_type'] == 'income' else 'expense' if record['transaction_type'] == 'expense' else 'neutral'
    if record['transaction_type'] == 'income':
        receipt_kind = _clean(row.get('receipt_kind') or 'other_income').lower()
        if receipt_kind not in RECEIPT_KINDS:
            record.update(decision='review', reason='收款类型无法识别')
        else:
            record['receipt_kind'] = receipt_kind
            if receipt_kind == 'internal_transfer':
                record['transaction_type'] = 'transfer'
    return record


_CATEGORY_KEYWORDS = [
    ('餐饮', ['美团', '饿了么', '麦当劳', '肯德基', 'KFC', '星巴克', '瑞幸', '蜜雪冰城', '海底捞',
              '餐饮', '食堂', '饭店', '火锅', '烧烤', '奶茶', '咖啡', '外卖', '小吃', '面馆',
              '粥', '汉堡', '炸鸡', '披萨', '便当', '快餐', '烘焙', '蛋糕', '面包', '便利店',
              '全家', '罗森', '7-11', '711', '喜茶', '奈雪', '茶百道', '古茗', 'CoCo',
              '华莱士', '德克士', '必胜客', '呷哺', '杨国福', '张亮', '沙县', '兰州拉面',
              '超市', '永辉', '沃尔玛', '家乐福', '大润发', '盒马']),
    ('交通', ['滴滴', '高德', '地铁', '公交', '打车', '出租', '高铁', '火车', '铁路', '12306',
              '航空', '机票', '飞机', '机场', '加油', '中石化', '中石油', '停车', 'ETC',
              '单车', '哈啰', '青桔', '美团单车', '共享', '出行', '网约车', '曹操', '花小猪',
              '船票', '轮渡', '客运', '汽车站', '充电站', '蔚来', '特来电']),
    ('购物', ['淘宝', '天猫', '京东', '拼多多', '唯品会', '苏宁', '抖音商城', '小红书商城',
              '商城', '服饰', '服装', '鞋', '箱包', '皮具', '数码', '电器', '家电', '超市',
              '百货', '购物', '专卖店', '旗舰店', '专营店', '折扣', '奥莱', '屈臣氏',
              '名创优品', 'MINISO', '无印良品', 'MUJI', '宜家', '五金', '建材', '家具',
              '宠物', '花店', '母婴', '玩具', '图书', '文具', '眼镜', '手表', '珠宝',
              '化妆品', '护肤', '美妆', '丝芙兰', '雅诗兰黛', '欧莱雅', '完美日记',
              '得物', '闲鱼', '二手', '网购', '电商', '订单', '商城']),
    ('娱乐', ['电影', '影院', '万达影城', '游戏', 'Steam', '腾讯游戏', '网易游戏', 'B站',
              '哔哩哔哩', 'bilibili', '爱奇艺', '优酷', '腾讯视频', '芒果', 'Netflix',
              '网易云音乐', 'QQ音乐', '酷狗', '酷我', 'KTV', '酒吧', '演出', '门票',
              '旅游', '景点', '乐园', '迪士尼', '欢乐谷', '密室', '剧本杀', '桌游',
              '健身', '运动', '体育', '游泳', '瑜伽', '舞蹈', '台球', '保龄球',
              '网吧', '网咖', '棋牌', '休闲', '娱乐', '会员', '订阅', '充值',
              '爱奇艺', '优酷', '视频', '音乐', '播客', '直播', '打赏', '斗鱼', '虎牙',
              '书店', '阅读', '小说', '起点', '晋江', '微信读书', 'Kindle']),
    ('学习', ['学费', '教材', '课程', '培训', '教育', '考试', '报名', '书店', '图书',
              '文具', '学习', '网课', '考研', '考证', '英语', '四六级', '雅思', '托福',
              '计算机', '编程', 'Python', 'Java', ' Coursera', 'Udemy', '极客时间',
              '得到', '知识星球', '知乎', '学堂', '大学', '学院', '图书馆',
              '打印机', '打印', '复印', '快递', '论文', '查重', '知网', '万方']),
    ('房租', ['房租', '租金', '物业', '水费', '电费', '燃气', '暖气', '宽带', '网费',
              '租房', '公寓', '中介', '押金', '水电', '缴费', '国家电网', '水务',
              '燃气公司', '供暖', '物业费']),
    ('话费', ['话费', '手机充值', '中国移动', '中国联通', '中国电信', '流量充值',
              '话费充值', '充值话费', '套餐', 'SIM', '移动', '联通', '电信']),
    ('医疗', ['医院', '诊所', '药店', '药房', '体检', '挂号', '门诊', '医疗',
              '医保', '牙科', '眼科', '中医', '保健', '药品', '胶囊', '片',
              '大药房', '同仁堂', '老百姓', '益丰', '一心堂', '健康',
              '疫苗', '核酸', '检测', '康复', '理疗', '心理咨询']),
]


def _guess_category(merchant, activity='', note=''):
    text = f'{merchant} {activity} {note}'
    for category, keywords in _CATEGORY_KEYWORDS:
        for kw in keywords:
            if kw in text:
                return category
    return '未分类'


def _wechat(row, line_number):
    activity, direction, status = (_clean(row.get(key)) for key in ('交易类型', '收/支', '当前状态'))
    kind = ('refund' if '退款' in activity or '退款' in status
            else 'transfer' if any(label in activity for label in ('充值', '提现', '零钱通', '转入', '转出')) or direction == '不计收支'
            else 'expense' if direction == '支出' else 'income' if direction == '收入' else '')
    merchant_raw = row.get('交易对方')
    note_raw = row.get('备注') or row.get('商品')
    auto_category = _guess_category(merchant_raw, activity, note_raw) if kind == 'expense' else '未分类'
    record = _common(line_number, kind, row.get('金额(元)') or row.get('金额'),
                     row.get('交易时间'), auto_category, merchant_raw,
                     note_raw, row.get('交易单号'), '',
                     '退款成功' if kind == 'refund' and status == '已全额退款' else status)
    record['merchant_order_id'] = _clean(row.get('商户单号'))[:100]
    record['direction'] = direction
    record['related_source_id'] = _clean(row.get('原交易单号') or row.get('原订单交易号'))[:100]
    if kind == 'income':
        if '红包' in activity:
            record['receipt_kind'] = 'red_packet'
        elif '转账' in activity:
            record['receipt_kind'] = 'external_transfer'
        elif any(label in activity for label in ('收款', '工资', '奖学金', '生活费')):
            record['receipt_kind'] = 'other_income'
        elif record['decision'] == 'import':
            record.update(decision='review', reason='收款用途不明，请核对后入账')
    if kind == 'expense' and '转账' in activity:
        record['category'] = '转账支出'
    if not record['source_transaction_id']:
        record.update(decision='review', reason='缺少交易单号，无法跨文件去重')
    return record


def _batches(uid):
    return [entry for entry in learning_store.read_entries(uid) if entry.get('kind') == 'import_batch']


def active_rows(uid):
    rows, seen = [], set()
    for batch in sorted(_batches(uid), key=lambda entry: (entry.get('created_at', ''), entry['id'])):
        if batch.get('status') != 'committed':
            continue
        source = batch['source']
        for row in batch.get('rows', []):
            if row.get('decision', 'import') != 'import':
                continue
            source_id = row.get('source_transaction_id')
            key = (source, source_id) if source_id else None
            if key and key in seen:
                continue
            if key:
                seen.add(key)
            rows.append({**row, 'id': f"{batch['id']}:{row['row_id']}",
                         'source': source, 'batch_id': batch['id'],
                         'timestamp': batch.get('created_at')})
    return rows


def _manual_purchases(uid):
    from .user_data_service import UserDataService
    return UserDataService().manual_purchase_rows(uid)


def supplement_purchase(uid, payload):
    """An explicit original transaction number makes repeated backfill idempotent."""
    if not isinstance(payload, dict) or set(payload) - {'source', 'transaction_id', 'date', 'amount', 'category', 'merchant'}:
        raise ValueError('原支出内容无效')
    source = payload.get('source')
    identifier = payload.get('transaction_id')
    if source not in ('wechat', 'standard') or not isinstance(identifier, str) or not 0 < len(identifier.strip()) <= 100:
        raise ValueError('请选择来源并填写原交易号；没有交易号可先手动记账再选择关联')
    identifier = identifier.strip()
    amount, spent = _amount(payload.get('amount')), _date(payload.get('date'))
    category, merchant = payload.get('category'), payload.get('merchant', '')
    if not isinstance(category, str) or not category.strip() or len(category.strip()) > 50:
        raise ValueError('原支出类别无效')
    if not isinstance(merchant, str) or len(merchant.strip()) > 100:
        raise ValueError('商户名称过长')
    previous = next((row for row in active_rows(uid) if row.get('source') == source
                     and row.get('source_transaction_id') == identifier), None)
    if previous:
        if previous['transaction_type'] != 'expense' or Decimal(previous['amount']) != Decimal(amount) or previous['date'] != spent:
            raise ValueError('该原交易号已有不同记录，请选择已有交易或核对后修改')
        return {'source': source, 'source_transaction_id': identifier, 'date': spent,
                'amount': amount, 'category': previous['category'], 'merchant': previous.get('merchant', '')}
    digest = hashlib.sha256(f'original:{source}:{identifier}'.encode()).hexdigest()
    row = _common(1, 'expense', amount, spent, category, merchant, '用户补录原支出', identifier, '', 'completed')
    row['supplemental'] = True
    learning_store.write_entry(uid, f'import:{digest}', {'kind': 'import_batch', 'status': 'committed',
                               'source': source, 'file_sha256': digest, 'parser_version': PARSER_VERSION,
                               'supplemental': True, 'rows': [row]})
    return {'source': source, 'source_transaction_id': identifier, 'date': spent,
            'amount': amount, 'category': category.strip(), 'merchant': merchant.strip()}


def manual_has_refunds(uid, identifier):
    return any(row.get('transaction_type') == 'refund' and _related_key(row) == ('manual', str(identifier))
               for row in active_rows(uid))


def _related_key(row):
    return (row.get('related_source') or row.get('source'), row.get('related_source_id'))


def _purchase_options(source, prior, rows=(), manual=()):
    return [{'source_transaction_id': row['source_transaction_id'], 'date': row['date'],
             'amount': row['amount'], 'merchant': row.get('merchant', ''),
             'category': row.get('category', '未分类'), 'source': row.get('source', source)}
            for row in chain(prior, rows, manual) if row.get('transaction_type') == 'expense'
            and row.get('source_transaction_id')
            and row.get('decision', 'import') == 'import']


def _validate_refunds(rows, manual=()):
    rows = [row for row in rows if row.get('decision', 'import') == 'import']
    remaining = {(row.get('source'), row['source_transaction_id']): Decimal(str(row['amount']))
                 for row in chain(rows, manual) if row.get('transaction_type') == 'expense' and row.get('source_transaction_id')}
    for row in rows:
        if row.get('transaction_type') != 'refund':
            continue
        key = _related_key(row)
        if key not in remaining:
            raise ValueError('退款缺少可关联的原支出')
        remaining[key] -= Decimal(str(row['amount']))
        if remaining[key] < 0:
            raise ValueError('退款总额不能超过关联支出')


def _apply_custom(uid, row, option_id):
    from .bill_categories import resolve
    option = resolve(uid, option_id)
    original = row.get('original_transaction_type') or row['transaction_type']
    incoming = row.get('direction') in ('income', '收入', 'neutral')
    allowed = ({'expense', 'review'} if original == 'expense' else
               {'refund', 'review'} if original == 'refund' else
               {'income', 'refund', 'transfer', 'review'} if original == 'income' or incoming else
               {'expense', 'transfer', 'review'})
    kind = option['accounting_type']
    if kind not in allowed:
        raise ValueError('此分类的账务类型不适用于该行')
    if not row.get('original_transaction_type'):
        row['original_category'] = row.get('category')
        row['original_receipt_kind'] = row.get('receipt_kind')
    row.update(category_id=option['id'], category=option['name'], custom_category_name=option['name'],
               original_transaction_type=original, transaction_type=kind,
               receipt_kind='other_income' if kind == 'income' else 'internal_transfer' if kind == 'transfer' else None)


def _finish_rows(uid, source, rows, prior):
    manual = _manual_purchases(uid)
    purchases = {(row.get('source', source), row['source_transaction_id']): row
                 for row in chain(prior, rows, manual) if row.get('transaction_type') == 'expense'
                 and row.get('decision', 'import') == 'import' and row.get('source_transaction_id')}
    for row in rows:
        if row.get('decision') == 'deleted':
            continue
        if row['transaction_type'] == 'review':
            row.update(decision='review', reason='自定义待核对，不参与统计')
        elif row['transaction_type'] == 'refund':
            if not row.get('related_source_id'):
                row.update(decision='review', reason='待关联退款，不参与统计')
            else:
                key = _related_key({**row, 'source': source})
                if key not in purchases:
                    raise ValueError('所选退款缺少原支出，请重新核对')
                row.update(decision='import', reason='')
                row['category'] = purchases[key].get('category') or '未分类'
        elif row['transaction_type'] == 'income' and not row.get('receipt_kind'):
            row.update(decision='review', reason='自定义待核对，不参与统计')
        elif row.get('decision') == 'review' and row.get('reason') in (
                '自定义待核对，不参与统计', '待关联退款，不参与统计'):
            row.update(decision='import', reason='')
    _validate_refunds(prior + [{**row, 'source': source} for row in rows], manual)


def preview(uid, content, hint='auto', file_type='csv'):
    source, rows = _read(content, hint, file_type)
    digest = hashlib.sha256(content).hexdigest()
    existing_batch = learning_store.read_entry(uid, f'import:{digest}')
    prior = active_rows(uid)
    order_matches = {}
    for purchase in chain(prior, rows):
        if (purchase.get('transaction_type') == 'expense' and purchase.get('merchant_order_id')
                and purchase.get('source_transaction_id') and purchase.get('source', source) == source):
            order_matches.setdefault(purchase['merchant_order_id'], set()).add(purchase['source_transaction_id'])
    for row in rows:
        if row.get('transaction_type') == 'refund' and not row.get('related_source_id'):
            matches = order_matches.get(row.get('merchant_order_id'), set())
            if len(matches) == 1:
                row['related_source_id'] = next(iter(matches))
            elif row['decision'] == 'import':
                row.update(decision='review', reason='未找到唯一的原交易，请手动关联')
    categories = {(r['source'], r['source_transaction_id']): r['category'] for r in prior
                  if r.get('transaction_type') == 'expense' and r.get('source_transaction_id')}
    categories.update({(source, r['source_transaction_id']): r['category'] for r in rows
                       if r.get('transaction_type') == 'expense' and r.get('source_transaction_id')})
    identifiers = {(r['source'], r['source_transaction_id']) for r in prior if r.get('source_transaction_id')}
    purchase_ids = {(r['source'], r['source_transaction_id']) for r in prior
                    if r.get('transaction_type') == 'expense' and r.get('source_transaction_id')}
    purchase_ids.update((source, r['source_transaction_id']) for r in rows
                        if r.get('transaction_type') == 'expense' and r['decision'] == 'import'
                        and r.get('source_transaction_id'))
    in_file = set()
    for row in rows:
        key = (source, row.get('source_transaction_id'))
        if row.get('transaction_type') == 'refund' and row.get('related_source_id'):
            row['category'] = categories.get((source, row['related_source_id']), row['category'])
        if row['decision'] != 'import':
            continue
        if row.get('source_transaction_id') and (key in identifiers or key in in_file):
            row.update(decision='duplicate', reason='该来源交易号已入账或在文件中重复')
        elif row['transaction_type'] == 'refund' and (source, row.get('related_source_id')) not in purchase_ids:
            row.update(decision='review', reason='找不到可关联的原支出')
        if row.get('source_transaction_id'):
            in_file.add(key)
        if row['transaction_type'] == 'expense' and row['decision'] == 'import' and row.get('source_transaction_id'):
            purchase_ids.add(key)
    counts = {kind: sum(row['decision'] == kind for row in rows)
              for kind in ('import', 'duplicate', 'review')}
    by_type = {kind: sum(row['decision'] == 'import' and row['transaction_type'] == kind for row in rows)
               for kind in VALID_TYPES}
    net_change = sum((Decimal(row['amount']) * (-1 if row['transaction_type'] == 'refund' else 1)
                      for row in rows if row['decision'] == 'import'
                      and row['transaction_type'] in ('expense', 'refund')), Decimal('0'))
    return {'source': source, 'sha256': digest, 'parser_version': PARSER_VERSION,
            'rows': rows, 'counts': counts, 'by_type': by_type, 'net_change': str(net_change),
            'purchase_options': _purchase_options(source, prior, rows, _manual_purchases(uid)),
            'already_imported': bool(existing_batch and existing_batch.get('status') == 'committed')}


def confirm(uid, content, expected_digest, selected_ids, categories, hint='auto', file_type='csv', edits=None):
    result = preview(uid, content, hint, file_type)
    if expected_digest != result['sha256']:
        raise ValueError('文件已变化，请重新预览')
    if not isinstance(selected_ids, list) or len(selected_ids) != len(set(selected_ids)) or not all(type(i) is int for i in selected_ids):
        raise ValueError('选择的行号无效')
    if not isinstance(categories, dict):
        raise ValueError('类别修改无效')
    edits = edits or {}
    if not isinstance(edits, dict):
        raise ValueError('交易修改无效')
    existing_batch = learning_store.read_entry(uid, f"import:{result['sha256']}") if result['already_imported'] else None
    existing_row_ids = {row['row_id'] for row in existing_batch['rows']
                        if row.get('decision', 'import') == 'import'} if existing_batch else set()
    selected_ids = [row_id for row_id in selected_ids if row_id not in existing_row_ids]
    if existing_batch and not selected_ids:
        return existing_batch, False
    by_id = {row['row_id']: row for row in result['rows']}
    chosen = []
    for row_id in selected_ids:
        row = by_id.get(row_id)
        change = edits.get(str(row_id), {})
        if not isinstance(change, dict) or set(change) - {'related_source_id', 'related_source', 'receipt_kind', 'category_id'}:
            raise ValueError('交易修改无效')
        can_link = (row and row.get('transaction_type') == 'refund' and row.get('decision') == 'review'
                    and row.get('reason') in ('未找到唯一的原交易，请手动关联', '找不到可关联的原支出')
                    and change.get('related_source_id'))
        can_classify = (row and row.get('transaction_type') == 'income' and row.get('decision') == 'review'
                        and row.get('reason') == '收款用途不明，请核对后入账'
                        and (change.get('receipt_kind') in RECEIPT_KINDS or change.get('category_id')))
        can_defer = (row and row.get('decision') == 'review' and row.get('reason') in (
                    '未找到唯一的原交易，请手动关联', '找不到可关联的原支出') and change.get('category_id'))
        if not row or (row['decision'] != 'import' and not can_link and not can_classify and not can_defer):
            raise ValueError('存在不可导入的行，请重新预览')
        category = categories.get(str(row_id), row['category'])
        if not isinstance(category, str) or not category.strip() or len(category.strip()) > 50:
            raise ValueError('消费类别无效')
        selected = {**row, 'category': category.strip(), 'source': result['source'], 'decision': 'import', 'reason': ''}
        if change.get('category_id'):
            _apply_custom(uid, selected, change['category_id'])
        if 'related_source_id' in change:
            if selected['transaction_type'] != 'refund' or not isinstance(change['related_source_id'], str):
                raise ValueError('退款关联无效')
            selected['related_source_id'] = _clean(change['related_source_id'])[:100]
        if 'related_source' in change:
            if change['related_source'] not in ('manual', 'wechat', 'standard'):
                raise ValueError('原支出来源无效')
            selected['related_source'] = change['related_source']
        if 'receipt_kind' in change:
            kind = change['receipt_kind']
            if (selected['transaction_type'] not in ('income', 'transfer') or kind not in RECEIPT_KINDS
                    or (selected['transaction_type'] == 'transfer' and selected.get('direction') not in ('income', '收入', 'neutral'))):
                raise ValueError('收款类型无效')
            selected['receipt_kind'] = kind
            selected['transaction_type'] = 'transfer' if kind == 'internal_transfer' else 'income'
            selected.pop('category_id', None)
            selected.pop('custom_category_name', None)
        chosen.append(selected)
    _finish_rows(uid, result['source'], chosen, active_rows(uid))
    if not chosen:
        raise ValueError('请选择至少一行可导入的记录')
    if len(json.dumps(chosen, ensure_ascii=False).encode('utf-8')) > MAX_BATCH_BYTES:
        raise ValueError('所选记录过多，请按日期拆分账单后导入')
    payload = ({k: v for k, v in existing_batch.items() if k not in ('id', 'created_at', 'updated_at')}
               if existing_batch else {'kind': 'import_batch', 'status': 'committed',
                                       'source': result['source'], 'file_sha256': result['sha256'],
                                       'parser_version': PARSER_VERSION})
    replacing = {row['row_id'] for row in chosen}
    payload['rows'] = ([row for row in existing_batch['rows'] if row['row_id'] not in replacing]
                       if existing_batch else []) + chosen
    if len(json.dumps(payload, ensure_ascii=False).encode('utf-8')) > MAX_BATCH_BYTES:
        raise ValueError('累计导入记录过多，请按日期拆分账单')
    record = learning_store.write_entry(uid, f"import:{result['sha256']}", payload)
    return record, True


def undo(uid, digest):
    batch = learning_store.read_entry(uid, f'import:{digest}')
    if not batch or batch.get('kind') != 'import_batch':
        raise ValueError('导入批次不存在')
    if batch.get('status') == 'cancelled':
        return batch
    own_ids = {row.get('source_transaction_id') for row in batch['rows']
               if row['transaction_type'] == 'expense' and row.get('source_transaction_id')}
    for other in active_rows(uid):
        if (other['batch_id'] != batch['id'] and other.get('transaction_type') == 'refund'
                and other.get('related_source_id') in own_ids
                and _related_key(other)[0] == batch['source']):
            raise ValueError('其他批次的退款仍关联这份账单，暂不能撤销')
    return learning_store.write_entry(uid, batch['id'], {k: v for k, v in batch.items()
                                                        if k not in ('id', 'created_at', 'updated_at')} | {'status': 'cancelled'})


def batch_detail(uid, digest):
    batch = learning_store.read_entry(uid, f'import:{digest}')
    if not batch or batch.get('kind') != 'import_batch':
        raise ValueError('导入批次不存在')
    purchases = _purchase_options(batch['source'], active_rows(uid), manual=_manual_purchases(uid))
    return {'sha256': digest, 'source': batch['source'], 'status': batch['status'],
            'rows': batch.get('rows', []), 'purchase_options': purchases,
            'can_revert_edit': bool(batch.get('edit_history'))}


def _write_revised(uid, batch, rows, history):
    others = [row for row in active_rows(uid) if row['batch_id'] != batch['id']]
    _finish_rows(uid, batch['source'], rows, others)
    record = {k: v for k, v in batch.items() if k not in ('id', 'created_at', 'updated_at')}
    record.update(rows=rows, edit_history=history)
    if len(json.dumps(record, ensure_ascii=False).encode('utf-8')) > MAX_BATCH_BYTES:
        raise ValueError('修改记录过多，请先撤销整批后重新导入')
    return learning_store.write_entry(uid, batch['id'], record)


def remove_row(uid, digest, row_id):
    """Remove one imported ledger entry, keeping its source and undo history."""
    batch = learning_store.read_entry(uid, f'import:{digest}')
    if not batch or batch.get('kind') != 'import_batch' or batch.get('status') != 'committed':
        raise ValueError('导入批次不存在或已撤销')
    rows = [dict(row) for row in batch['rows']]
    target = next((row for row in rows if row['row_id'] == row_id), None)
    if target is None:
        raise ValueError('交易不存在')
    if target.get('decision') == 'deleted':
        return batch
    if target.get('transaction_type') == 'expense':
        key = (batch['source'], target.get('source_transaction_id'))
        if any(row.get('transaction_type') == 'refund' and _related_key(row) == key
               for row in active_rows(uid)):
            raise ValueError('退款仍关联这笔支出，请先删除退款或修改退款关联')
    previous = {key: target.get(key) for key in ('decision', 'reason')}
    target.update(decision='deleted', reason='用户删除，不参与统计')
    history = (batch.get('edit_history') or [])[-19:] + [{'row_id': row_id, 'previous': previous}]
    return _write_revised(uid, batch, rows, history)


def update_row(uid, digest, row_id, changes):
    batch = learning_store.read_entry(uid, f'import:{digest}')
    if not batch or batch.get('kind') != 'import_batch' or batch.get('status') != 'committed':
        raise ValueError('导入批次不存在或已撤销')
    if not isinstance(changes, dict) or not changes or set(changes) - {'category', 'receipt_kind', 'related_source_id', 'related_source', 'category_id'}:
        raise ValueError('修改内容无效')
    rows = [dict(row) for row in batch['rows']]
    target = next((row for row in rows if row['row_id'] == row_id), None)
    if target is None:
        raise ValueError('交易不存在')
    if target.get('decision') == 'deleted':
        raise ValueError('交易已删除，请先撤销上次修改后再编辑')
    previous = {key: target.get(key) for key in ('category', 'category_id', 'receipt_kind',
                'related_source_id', 'related_source', 'transaction_type', 'original_transaction_type',
                'custom_category_name', 'original_category', 'original_receipt_kind', 'decision', 'reason')}
    if changes.get('category_id'):
        _apply_custom(uid, target, changes['category_id'])
    elif 'category_id' in changes:
        target.pop('category_id', None)
        target.pop('custom_category_name', None)
        if target.get('original_transaction_type'):
            target['transaction_type'] = target['original_transaction_type']
            target['receipt_kind'] = target.get('original_receipt_kind')
            target['category'] = target.get('original_category') or '未分类'
    if 'category' in changes:
        value = changes['category']
        if target['transaction_type'] != 'expense' or not isinstance(value, str) or not value.strip() or len(value.strip()) > 50:
            raise ValueError('消费类别无效')
        target['category'] = value.strip()
        target.pop('category_id', None)
        target.pop('custom_category_name', None)
    if 'receipt_kind' in changes:
        value = changes['receipt_kind']
        if ((target.get('original_transaction_type') or target['transaction_type']) not in ('income', 'transfer') or value not in RECEIPT_KINDS
                or (target['transaction_type'] == 'transfer' and target.get('direction') not in ('income', '收入', 'neutral'))):
            raise ValueError('收款类型无效')
        target['receipt_kind'] = value
        target['transaction_type'] = 'transfer' if value == 'internal_transfer' else 'income'
        target.pop('category_id', None)
        target.pop('custom_category_name', None)
    if 'related_source_id' in changes:
        value = changes['related_source_id']
        if target['transaction_type'] != 'refund' or not isinstance(value, str) or len(value.strip()) > 100:
            raise ValueError('退款关联无效')
        target['related_source_id'] = value.strip()
    if 'related_source' in changes:
        if changes['related_source'] not in ('manual', 'wechat', 'standard'):
            raise ValueError('原支出来源无效')
        target['related_source'] = changes['related_source']
    if previous == {key: target.get(key) for key in previous}:
        return batch
    history = (batch.get('edit_history') or [])[-19:] + [{'row_id': row_id, 'previous': previous}]
    return _write_revised(uid, batch, rows, history)


def revert_last_edit(uid, digest):
    batch = learning_store.read_entry(uid, f'import:{digest}')
    if not batch or batch.get('kind') != 'import_batch' or batch.get('status') != 'committed':
        raise ValueError('导入批次不存在或已撤销')
    history = list(batch.get('edit_history') or [])
    if not history:
        raise ValueError('没有可撤销的逐笔修改')
    last = history.pop()
    rows = [dict(row) for row in batch['rows']]
    target = next(row for row in rows if row['row_id'] == last['row_id'])
    target.update(last['previous'])
    return _write_revised(uid, batch, rows, history)


def combined_transactions(uid, manual_rows, limit):
    manual = [{**row, 'transaction_type': row.get('transaction_type') or 'expense',
               'source': row.get('source') or 'manual',
               'source_transaction_id': row.get('source_transaction_id') or str(row.get('id'))} for row in manual_rows]
    rows = manual + active_rows(uid)
    purchases = {(row.get('source'), row.get('source_transaction_id')): row
                 for row in rows if row.get('transaction_type') == 'expense' and row.get('source_transaction_id')}
    refunds = {}
    for row in rows:
        if row.get('transaction_type') == 'refund':
            key = _related_key(row)
            refunds[key] = refunds.get(key, Decimal('0')) + Decimal(str(row['amount']))
            if key in purchases:
                row['category'] = purchases[key].get('category', '未分类')
    for row in rows:
        if row.get('transaction_type') == 'expense':
            amount = Decimal(str(row['amount']))
            refunded = refunds.get((row.get('source'), row.get('source_transaction_id')), Decimal('0'))
            row.update(original_amount=str(amount), refunded_amount=str(refunded), net_amount=str(amount - refunded))
    rows.sort(key=lambda row: (str(row.get('date') or row.get('spent_at') or '')[:10],
                               str(row.get('timestamp') or '')), reverse=True)
    return rows[:limit]


def imported_spending(uid, start, end):
    total = Decimal('0')
    for row in active_rows(uid):
        if start <= str(row.get('date', ''))[:10] <= end:
            amount = Decimal(str(row['amount']))
            if row['transaction_type'] == 'expense':
                total += amount
            elif row['transaction_type'] == 'refund':
                total -= amount
    return total


def imported_cashflow(uid, start, end):
    gross = refunds = income = Decimal('0')
    breakdown = {'red_packet': Decimal('0'), 'external_transfer': Decimal('0'), 'other_income': Decimal('0')}
    counts = {'expense': 0, 'refund': 0, 'income': 0}
    for row in active_rows(uid):
        if not start <= str(row.get('date', ''))[:10] <= end:
            continue
        amount = Decimal(str(row['amount']))
        kind = row['transaction_type']
        if kind == 'expense':
            gross += amount
            counts['expense'] += 1
        elif kind == 'refund':
            refunds += amount
            counts['refund'] += 1
        elif kind == 'income':
            receipt = row.get('receipt_kind') or 'other_income'
            if receipt not in breakdown:
                raise ValueError('收款分类无效')
            income += amount
            breakdown[receipt] += amount
            counts['income'] += 1
    return {'gross_expense': gross, 'refunds': refunds, 'spent': gross - refunds,
            'income': income, 'income_breakdown': breakdown, 'counts': counts}
